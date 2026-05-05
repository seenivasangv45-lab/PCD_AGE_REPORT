"""
Monthly Report Dashboard – Web Application
============================================
Flask-based multi-user web app that processes PCD_20 + Age_24 data
and generates interactive dashboard reports.

Supports:
  - Two separate file uploads (PCD_20 and Age_24)
  - Both CSV and XLSX formats
  - Auto-detection by filename or column structure
  - Multiple simultaneous users

Double-click run_server.bat (Windows) or run_server.sh (Mac/Linux) to launch.
"""

import os
import json
import uuid
import socket
import threading
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import Flask, request, jsonify, send_file, send_from_directory

# ══════════════════════════════════════════════════════════════════════════════
#  CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════

SELF_PAY_VALUES = {"1-Self Pay"}
WC_MVA_VALUES   = {"30-WC Non-Responsive", "23-Auto", "20-Work Comp"}
EPS_VALUES      = {"9-EPS"}

DASHBOARD_ROW_MAP = {
    "ins_charges": 2, "ins_visits": 3, "ins_avg_charge": 4,
    "ins_paid_closed": 5, "ins_no_payment": 6, "ins_partial": 7,
    "ins_patient_balance": 8, "ins_credit_balance": 9,
    "ins_primary_paid_pct": 10, "ins_net_collection": 11,
    "ins_contractual_adj": 12, "ins_avg_collection": 13,
    "sp_charges": 15, "sp_visits": 16, "sp_avg_charge": 17,
    "sp_net_collection": 18, "sp_contractual_adj": 19,
    "eps_charges": 21, "eps_visits": 22, "eps_avg_charge": 23,
    "eps_net_collection": 24, "eps_contractual_adj": 25,
    "wc_charges": 27, "wc_visits": 28, "wc_avg_charge": 29,
    "wc_net_collection": 30, "wc_contractual_adj": 31,
    "total_charges": 32, "total_visits": 33, "total_payments": 34,
    "total_adjustments": 35, "gcr_pct": 36, "ncr_pct": 37,
}

ROW_LABELS = {
    2: "Charges", 3: "# Visit", 4: "Avg Charge/Visit",
    5: "Paid & Closed", 6: "No Payment Visit", 7: "Partially Paid",
    8: "Patient Balance", 9: "Credit Balance",
    10: "Primary Paid %", 11: "Net Collection",
    12: "Contractual Adj", 13: "Avg Collection/Visit",
    15: "Charges", 16: "# Visit", 17: "Avg Charge/Visit",
    18: "Net Collection", 19: "Contractual Adj",
    21: "Charges", 22: "# Visit", 23: "Avg Charge/Visit",
    24: "Net Collection", 25: "Contractual Adj",
    27: "Charges", 28: "# Visit", 29: "Avg Charge/Visit",
    30: "Net Collection", 31: "Contractual Adj",
    32: "Total Charges", 33: "Total Visits", 34: "Total Payments",
    35: "Total Adjustments", 36: "GCR%", 37: "NCR%",
}

MONEY_FMT = '$#,##0.00'
PCT_FMT   = '0.00%'
INT_FMT   = '#,##0'
AVG_FMT   = '$#,##0.00'

FORMAT_MAP = {
    "ins_charges": MONEY_FMT, "ins_visits": INT_FMT, "ins_avg_charge": AVG_FMT,
    "ins_paid_closed": INT_FMT, "ins_no_payment": INT_FMT, "ins_partial": INT_FMT,
    "ins_patient_balance": INT_FMT, "ins_credit_balance": INT_FMT,
    "ins_primary_paid_pct": PCT_FMT, "ins_net_collection": MONEY_FMT,
    "ins_contractual_adj": MONEY_FMT, "ins_avg_collection": AVG_FMT,
    "sp_charges": MONEY_FMT, "sp_visits": INT_FMT, "sp_avg_charge": AVG_FMT,
    "sp_net_collection": MONEY_FMT, "sp_contractual_adj": MONEY_FMT,
    "eps_charges": MONEY_FMT, "eps_visits": INT_FMT, "eps_avg_charge": AVG_FMT,
    "eps_net_collection": MONEY_FMT, "eps_contractual_adj": MONEY_FMT,
    "wc_charges": MONEY_FMT, "wc_visits": INT_FMT, "wc_avg_charge": AVG_FMT,
    "wc_net_collection": MONEY_FMT, "wc_contractual_adj": MONEY_FMT,
    "total_charges": MONEY_FMT, "total_visits": INT_FMT,
    "total_payments": MONEY_FMT, "total_adjustments": MONEY_FMT,
    "gcr_pct": PCT_FMT, "ncr_pct": PCT_FMT,
}

# Excel styles
_THIN   = Side(style='thin', color='B0C4DE')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEADER_FILL  = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT  = Font(bold=True, color='FFFFFF', name='Arial', size=10)
INS_HDR_FILL = PatternFill('solid', fgColor='1F4E79')
BODY_FONT    = Font(name='Arial', size=10)
BODY_FONT_B  = Font(name='Arial', size=10, bold=True)
AL_LEFT      = Alignment(horizontal='left', vertical='center')
AL_RIGHT     = Alignment(horizontal='right', vertical='center')
AL_CENTER    = Alignment(horizontal='center', vertical='center')
INS_FILL     = PatternFill('solid', fgColor='D6E4F0')
SP_FILL      = PatternFill('solid', fgColor='D6E4F0')
EPS_FILL     = PatternFill('solid', fgColor='E2EFDA')
WC_FILL      = PatternFill('solid', fgColor='FFF2CC')
TOTAL_FILL   = PatternFill('solid', fgColor='F2F2F2')
WHITE_FILL   = PatternFill('solid', fgColor='FFFFFF')

SECTION_ROW_FILLS = {}
for r in range(2, 14):  SECTION_ROW_FILLS[r] = INS_FILL
for r in range(15, 20): SECTION_ROW_FILLS[r] = SP_FILL
for r in range(21, 26): SECTION_ROW_FILLS[r] = EPS_FILL
for r in range(27, 32): SECTION_ROW_FILLS[r] = WC_FILL
for r in range(32, 38): SECTION_ROW_FILLS[r] = TOTAL_FILL

SECTION_HEADERS = {
    14: ("Self Pay", "4472C4"),
    20: ("EPS", "548235"),
    26: ("WC & MVA", "BF8F00"),
}

# ══════════════════════════════════════════════════════════════════════════════
#  FILE IDENTIFICATION — Auto-detect PCD_20 vs Age_24
# ══════════════════════════════════════════════════════════════════════════════

PCD20_REQUIRED = {"Inv_Num", "Crg_Amt", "Paid_Amt", "Adj_Amt", "Balance", "Svc_Date", "textbox13"}
AGE24_REQUIRED = {"Inv_Num", "Financial_Class", "textbox18"}


def detect_file_type(filepath: str) -> str:
    """
    Auto-detect whether a file is PCD_20 or Age_24.
    Detection priority:
      1. Filename contains 'pcd' or 'age' (case-insensitive)
      2. Column structure matching
    Returns: 'pcd' or 'age'
    """
    fname = Path(filepath).stem.lower().replace(" ", "").replace("_", "").replace("-", "")

    # Strategy 1: filename-based
    if "pcd" in fname and "20" in fname:
        return "pcd"
    if "age" in fname and "24" in fname:
        return "age"
    if "pcd" in fname:
        return "pcd"
    if "age" in fname:
        return "age"

    # Strategy 2: column-structure detection
    ext = Path(filepath).suffix.lower()
    try:
        if ext == '.csv':
            df = pd.read_csv(filepath, nrows=5)
        else:
            xls = pd.ExcelFile(filepath)
            # Check sheet names first
            for sn in xls.sheet_names:
                sn_clean = sn.lower().replace(" ", "").replace("_", "")
                if "pcd" in sn_clean:
                    return "pcd"
                if "age" in sn_clean:
                    return "age"
            df = pd.read_excel(filepath, sheet_name=0, nrows=5)

        cols = set(df.columns.astype(str))
        if PCD20_REQUIRED.issubset(cols):
            return "pcd"
        if AGE24_REQUIRED.issubset(cols):
            return "age"
        if "Svc_Date" in cols and "Crg_Amt" in cols:
            return "pcd"
        if "Financial_Class" in cols and "textbox18" in cols:
            return "age"
    except Exception:
        pass

    raise ValueError(
        f"Cannot determine file type for '{Path(filepath).name}'. "
        "Ensure filenames contain 'PCD' or 'Age', or the file has the "
        "expected columns (Svc_Date, Crg_Amt for PCD_20; "
        "Financial_Class, textbox18 for Age_24)."
    )


def get_sheet_for_type(filepath: str, file_type: str) -> str:
    """For XLSX files, find the right sheet. For CSV returns None."""
    ext = Path(filepath).suffix.lower()
    if ext == '.csv':
        return None
    xls = pd.ExcelFile(filepath)
    if len(xls.sheet_names) == 1:
        return xls.sheet_names[0]
    target = "pcd" if file_type == "pcd" else "age"
    for sn in xls.sheet_names:
        if target in sn.lower().replace(" ", "").replace("_", ""):
            return sn
    required = PCD20_REQUIRED if file_type == "pcd" else AGE24_REQUIRED
    for sn in xls.sheet_names:
        try:
            df = pd.read_excel(filepath, sheet_name=sn, nrows=5)
            if required.issubset(set(df.columns.astype(str))):
                return sn
        except Exception:
            continue
    return xls.sheet_names[0]


def load_single_file(filepath: str, file_type: str) -> pd.DataFrame:
    """Load a single PCD or Age file (CSV or XLSX) into a DataFrame."""
    ext = Path(filepath).suffix.lower()
    if ext == '.csv':
        df = pd.read_csv(filepath, dtype={"Inv_Num": str})
    else:
        sheet = get_sheet_for_type(filepath, file_type)
        df = pd.read_excel(filepath, sheet_name=sheet, dtype={"Inv_Num": str})
    df = df.loc[:, ~df.columns.str.match(r"^Unnamed:")]
    return df


# ══════════════════════════════════════════════════════════════════════════════
#  DATA PROCESSING
# ══════════════════════════════════════════════════════════════════════════════

def filter_crg_amt(df):
    return df[df["Crg_Amt"] >= 0].copy()

def normalize_signs(df):
    df[["Paid_Amt", "Adj_Amt"]] = df[["Paid_Amt", "Adj_Amt"]] * -1
    return df

def enrich_age24(age: pd.DataFrame) -> tuple:
    if "Financial_Class" not in age.columns or "textbox18" not in age.columns:
        age = age.copy()
        age["New_Fin_Class"] = ""
        age["SP_Lookup_Status"] = ""
        return age, set()
    age = age.copy()
    age["Inv_Num"] = age["Inv_Num"].astype(str).str.strip()
    age["Financial_Class"] = age["Financial_Class"].astype(str).str.strip()
    age["textbox18"] = pd.to_numeric(age["textbox18"], errors="coerce").fillna(0)
    age["New_Fin_Class"] = np.where(
        age["Financial_Class"].isin(SELF_PAY_VALUES), "Self", "Insurance")
    age["SP_Lookup_Status"] = ""
    is_self = age["New_Fin_Class"] == "Self"
    self_inv_nums = set(age.loc[is_self, "Inv_Num"])
    if not self_inv_nums:
        return age, set()
    mask_relevant = age["Inv_Num"].isin(self_inv_nums)
    relevant = age.loc[mask_relevant].copy()
    self_max = (relevant.loc[relevant["New_Fin_Class"] == "Self"]
                .groupby("Inv_Num")["textbox18"].max().rename("self_tb18"))
    ins_rows = relevant.loc[relevant["New_Fin_Class"] == "Insurance"]
    ins_max = (ins_rows.groupby("Inv_Num")["textbox18"].max().rename("ins_tb18")
               if len(ins_rows) > 0
               else pd.Series(dtype=float, name="ins_tb18"))
    status_df = self_max.to_frame().join(ins_max, how="left")
    conditions = [
        status_df["self_tb18"] <= 0,
        status_df["ins_tb18"].isna(),
        status_df["self_tb18"] > status_df["ins_tb18"],
    ]
    status_df["status"] = np.select(conditions,
                                    ["Not Eligible", "Eligible", "Eligible"],
                                    default="Not Eligible")
    status_map = status_df["status"].to_dict()
    age.loc[mask_relevant, "SP_Lookup_Status"] = (
        age.loc[mask_relevant, "Inv_Num"].map(status_map))
    eligible_set = set(status_df.loc[status_df["status"] == "Eligible"].index)
    return age, eligible_set

def add_category(df):
    raw = df["textbox13"].fillna("").astype(str).str.strip()
    conditions = [raw.isin(SELF_PAY_VALUES), raw.isin(WC_MVA_VALUES), raw.isin(EPS_VALUES)]
    df["Category"] = np.select(conditions, ["Self Pay", "WC & MVA", "EPS"], default="Insurance")
    return df

def group_by_inv_num(df):
    return (df.groupby("Inv_Num", as_index=False)
              .agg(g_Crg=("Crg_Amt", "sum"), g_Adj=("Adj_Amt", "sum"),
                   g_Paid=("Paid_Amt", "sum"), g_Bal=("Balance", "sum")))

def step5_assign_status(df, grouped, eligible_sp: set):
    g = grouped.set_index("Inv_Num")
    crg = g["g_Crg"].round(4); adj = g["g_Adj"].round(4)
    paid = g["g_Paid"].round(4); bal = g["g_Bal"].round(4)
    is_sp = g.index.isin(eligible_sp)
    conditions = [crg == 0, crg == adj, bal < 0,
                  (paid > 0) & (bal == 0), (bal > 0) & is_sp, crg == bal]
    choices = ["Zero Charges", "Fully adjusted", "Credit balance",
               "Paid & Closed", "Patient balance", "No payment visit"]
    status_series = pd.Series(np.select(conditions, choices, default="Partially Paid"),
                              index=g.index)
    df["Status"] = df["Inv_Num"].map(status_series)
    return df

def compute_dashboard_metrics(pcd: pd.DataFrame, target_date: datetime) -> dict:
    month_start = target_date.replace(day=1)
    month_end = (month_start.replace(year=month_start.year + 1, month=1)
                 if month_start.month == 12
                 else month_start.replace(month=month_start.month + 1))
    mth = pcd[(pcd["Svc_Date"] >= month_start) & (pcd["Svc_Date"] < month_end)]
    mth_nz = mth[mth["Crg_Amt"] != 0]
    ins = mth_nz[mth_nz["Category"] == "Insurance"]
    sp  = mth_nz[mth_nz["Category"] == "Self Pay"]
    eps = mth_nz[mth_nz["Category"] == "EPS"]
    wc  = mth_nz[mth_nz["Category"] == "WC & MVA"]

    def chg(s): return round(float(s["Crg_Amt"].sum()), 2)
    def vis(s): return int(s["Inv_Num"].nunique())
    def net(s): return round(float(s["Paid_Amt"].sum()), 2)
    def adj(s): return round(float(s["Adj_Amt"].sum()), 2)
    def dc(cat, status):
        cr = mth[mth["Category"] == cat]
        return int(cr.loc[cr["Status"] == status, "Inv_Num"].nunique())

    m = {}
    ic, iv, ip, ia = chg(ins), vis(ins), net(ins), adj(ins)
    ipc = dc("Insurance", "Paid & Closed")
    ifa = dc("Insurance", "Fully adjusted")
    inp_ = dc("Insurance", "No payment visit") + ifa
    ipt = dc("Insurance", "Partially Paid")
    ipb = dc("Insurance", "Patient balance")
    icb = dc("Insurance", "Credit balance")
    iavg = round(ic / iv, 10) if iv else 0.0
    ipg = ipc + ipt + ipb + icb
    ipp = round(ipg / iv, 10) if iv else 0.0
    iac = round(ip / ipg, 10) if ipg else 0.0
    m.update({"ins_charges": ic, "ins_visits": iv, "ins_avg_charge": iavg,
              "ins_paid_closed": ipc, "ins_no_payment": inp_, "ins_partial": ipt,
              "ins_patient_balance": ipb, "ins_credit_balance": icb,
              "ins_primary_paid_pct": ipp, "ins_net_collection": ip,
              "ins_contractual_adj": ia, "ins_avg_collection": iac})

    sc, sv = chg(sp), vis(sp)
    m.update({"sp_charges": sc, "sp_visits": sv,
              "sp_avg_charge": round(sc / sv, 10) if sv else 0.0,
              "sp_net_collection": net(sp), "sp_contractual_adj": adj(sp)})

    ec, ev = chg(eps), vis(eps)
    m.update({"eps_charges": ec, "eps_visits": ev,
              "eps_avg_charge": round(ec / ev, 10) if ev else 0.0,
              "eps_net_collection": net(eps), "eps_contractual_adj": adj(eps)})

    wc_, wv = chg(wc), vis(wc)
    m.update({"wc_charges": wc_, "wc_visits": wv,
              "wc_avg_charge": round(wc_ / wv, 10) if wv else 0.0,
              "wc_net_collection": net(wc), "wc_contractual_adj": adj(wc)})

    tc = ic + sc + ec + wc_
    tv = iv + sv + ev + wv
    tp = ip + m["sp_net_collection"] + m["eps_net_collection"] + m["wc_net_collection"]
    ta = ia + m["sp_contractual_adj"] + m["eps_contractual_adj"] + m["wc_contractual_adj"]
    gcr = round(tp / tc, 10) if tc else 0.0
    ncr = round((tp + ta) / tc, 10) if tc else 0.0
    m.update({"total_charges": round(tc, 2), "total_visits": tv,
              "total_payments": round(tp, 2), "total_adjustments": round(ta, 2),
              "gcr_pct": gcr, "ncr_pct": ncr})
    return m

def _generate_month_range(start: datetime, end: datetime) -> list:
    months = []
    current = start.replace(day=1)
    end_norm = end.replace(day=1)
    while current <= end_norm:
        months.append(current)
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1)
        else:
            current = current.replace(month=current.month + 1)
    return months


# ══════════════════════════════════════════════════════════════════════════════
#  EXCEL DASHBOARD GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def create_dashboard(wb, target_months):
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)
    num_months = len(target_months)
    last_col = 1 + num_months
    c = ws.cell(row=1, column=1, value="Insurance")
    c.fill = INS_HDR_FILL; c.font = HEADER_FONT; c.alignment = AL_CENTER; c.border = _BORDER
    for ci, dt in enumerate(target_months, 2):
        c = ws.cell(row=1, column=ci, value=dt)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = AL_CENTER
        c.border = _BORDER; c.number_format = 'MMM YYYY'
    for row_idx, label in ROW_LABELS.items():
        c = ws.cell(row=row_idx, column=1, value=label)
        fill = SECTION_ROW_FILLS.get(row_idx, WHITE_FILL)
        c.fill = fill; c.font = BODY_FONT_B if row_idx >= 32 else BODY_FONT
        c.alignment = AL_LEFT; c.border = _BORDER
    for row_idx, (title, color_hex) in SECTION_HEADERS.items():
        banner_fill = PatternFill('solid', fgColor=color_hex)
        banner_font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
        for ci in range(1, last_col + 1):
            c = ws.cell(row=row_idx, column=ci)
            c.fill = banner_fill; c.font = banner_font
            c.alignment = AL_CENTER; c.border = _BORDER
        ws.cell(row=row_idx, column=1, value=title)
    for row_idx in range(2, 38):
        if row_idx in SECTION_HEADERS:
            continue
        fill = SECTION_ROW_FILLS.get(row_idx, WHITE_FILL)
        font = BODY_FONT_B if row_idx >= 32 else BODY_FONT
        for ci in range(2, last_col + 1):
            c = ws.cell(row=row_idx, column=ci)
            c.fill = fill; c.font = font; c.alignment = AL_RIGHT; c.border = _BORDER
    ws.column_dimensions['A'].width = 24
    for ci in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 16
    ws.freeze_panes = "B2"
    section_last_rows = [13, 19, 25, 31, 37]
    thick_bottom = Border(left=_THIN, right=_THIN, top=_THIN,
                          bottom=Side(style='medium', color='1F4E79'))
    for row_idx in section_last_rows:
        for ci in range(1, last_col + 1):
            ws.cell(row=row_idx, column=ci).border = thick_bottom

def find_dashboard_column(ws, target_date: datetime):
    month_start = target_date.replace(day=1)
    for cell in ws[1]:
        v = cell.value
        if v is None: continue
        if isinstance(v, datetime):
            if v.year == month_start.year and v.month == month_start.month:
                return cell.column
        elif hasattr(v, 'year'):
            if v.year == month_start.year and v.month == month_start.month:
                return cell.column
    return None

def fill_dashboard(ws, metrics: dict, target_col: int):
    for key, row_idx in DASHBOARD_ROW_MAP.items():
        val = metrics.get(key)
        if val is None: continue
        c = ws.cell(row=row_idx, column=target_col)
        c.value = val
        c.number_format = FORMAT_MAP.get(key, '#,##0.00')


# ══════════════════════════════════════════════════════════════════════════════
#  FLASK APP
# ══════════════════════════════════════════════════════════════════════════════

app = Flask(__name__, static_folder='static')

UPLOAD_FOLDER = Path(__file__).parent / 'uploads'
OUTPUT_FOLDER = Path(__file__).parent / 'outputs'
UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)

app.config['MAX_CONTENT_LENGTH'] = 200 * 1024 * 1024


@app.route('/')
def index():
    return send_from_directory('static', 'index.html')


@app.route('/api/upload', methods=['POST'])
def upload_file():
    """Upload a single file. Auto-detects PCD_20 vs Age_24. Accepts CSV & XLSX."""
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files['file']
    if not file.filename:
        return jsonify({"error": "No file selected"}), 400
    ext = Path(file.filename).suffix.lower()
    if ext not in ('.xlsx', '.xls', '.csv'):
        return jsonify({"error": "Unsupported format. Please upload .xlsx or .csv files only."}), 400

    session_id = request.form.get('session_id') or str(uuid.uuid4())[:8]
    saved_name = f"{session_id}_{file.filename}"
    saved_path = UPLOAD_FOLDER / saved_name
    file.save(str(saved_path))

    try:
        file_type = detect_file_type(str(saved_path))
        detected_label = "PCD_20" if file_type == "pcd" else "Age_24"

        result = {
            "session_id": session_id,
            "original_filename": file.filename,
            "saved_filename": saved_name,
            "detected_type": file_type,
            "detected_label": detected_label,
            "format": ext.lstrip('.').upper(),
        }

        if file_type == "pcd":
            df = load_single_file(str(saved_path), "pcd")
            df["Svc_Date"] = pd.to_datetime(df["Svc_Date"], errors="coerce")
            df = df.dropna(subset=["Svc_Date"])
            result["min_date"] = df["Svc_Date"].min().strftime("%Y-%m-%d")
            result["max_date"] = df["Svc_Date"].max().strftime("%Y-%m-%d")
            result["total_rows"] = len(df)
        else:
            df = load_single_file(str(saved_path), "age")
            result["total_rows"] = len(df)

        return jsonify(result)
    except Exception as e:
        saved_path.unlink(missing_ok=True)
        return jsonify({"error": str(e)}), 400


@app.route('/api/generate', methods=['POST'])
def generate_report():
    """Generate dashboard using both uploaded files."""
    data = request.json
    session_id = data.get("session_id")
    start_date_str = data.get("start_date")
    end_date_str = data.get("end_date")
    pcd_filename = data.get("pcd_filename")
    age_filename = data.get("age_filename")

    if not all([session_id, start_date_str, end_date_str, pcd_filename, age_filename]):
        return jsonify({"error": "Missing required fields. Upload both files and select dates."}), 400

    pcd_path = UPLOAD_FOLDER / pcd_filename
    age_path = UPLOAD_FOLDER / age_filename
    if not pcd_path.exists():
        return jsonify({"error": "PCD_20 file not found. Please re-upload."}), 404
    if not age_path.exists():
        return jsonify({"error": "Age_24 file not found. Please re-upload."}), 404

    try:
        start_date = datetime.strptime(start_date_str, "%Y-%m-%d")
        end_date = datetime.strptime(end_date_str, "%Y-%m-%d")
    except ValueError:
        return jsonify({"error": "Invalid date format."}), 400

    if start_date > end_date:
        start_date, end_date = end_date, start_date
    target_months = _generate_month_range(start_date, end_date)

    try:
        pcd = load_single_file(str(pcd_path), "pcd")
        age = load_single_file(str(age_path), "age")

        pcd = filter_crg_amt(pcd)
        pcd = normalize_signs(pcd)
        age, eligible_sp = enrich_age24(age)
        pcd = add_category(pcd)
        grouped = group_by_inv_num(pcd)
        pcd = step5_assign_status(pcd, grouped, eligible_sp)
        pcd["Svc_Date"] = pd.to_datetime(pcd["Svc_Date"], errors="coerce")

        all_metrics = {}
        for target_date in target_months:
            month_key = target_date.strftime("%b %Y")
            all_metrics[month_key] = compute_dashboard_metrics(pcd, target_date)

        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        create_dashboard(wb, target_months)
        ws_dash = wb["Dashboard"]
        for target_date in target_months:
            month_key = target_date.strftime("%b %Y")
            dash_col = find_dashboard_column(ws_dash, target_date)
            if dash_col:
                fill_dashboard(ws_dash, all_metrics[month_key], dash_col)

        output_name = f"Report_{session_id}_{start_date.strftime('%b%Y')}_to_{end_date.strftime('%b%Y')}.xlsx"
        output_path = OUTPUT_FOLDER / output_name
        wb.save(str(output_path))

        formatted = {}
        for month_key, metrics in all_metrics.items():
            formatted[month_key] = {}
            for key, val in metrics.items():
                if key.endswith("_pct"):
                    formatted[month_key][key] = {"raw": val, "display": f"{val:.2%}"}
                elif isinstance(val, float):
                    formatted[month_key][key] = {"raw": val, "display": f"${val:,.2f}"}
                else:
                    formatted[month_key][key] = {"raw": val, "display": f"{val:,}"}

        return jsonify({
            "success": True,
            "months": list(all_metrics.keys()),
            "metrics": formatted,
            "download_file": output_name,
            "num_months": len(target_months),
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route('/api/download/<filename>')
def download_file(filename):
    filepath = OUTPUT_FOLDER / filename
    if not filepath.exists():
        return jsonify({"error": "File not found"}), 404
    return send_file(str(filepath), as_attachment=True, download_name=filename)


def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"


if __name__ == '__main__':
    port = 5000
    local_ip = get_local_ip()
    print("\n" + "=" * 60)
    print("  Monthly Report Dashboard Server")
    print("=" * 60)
    print(f"\n  Local access:   http://localhost:{port}")
    print(f"  Network access: http://{local_ip}:{port}")
    print(f"\n  Share the network URL with other users on your WiFi")
    print("=" * 60 + "\n")
    app.run(host='0.0.0.0', port=port, debug=False, threaded=True)
